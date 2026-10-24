"""Site inspections and annual audits: templates, numbering, saving the checklist,
findings, signatures, completion, printing, access rules and the Excel export."""
import json
import os
from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from app.extensions import db
from app.models import Inspection
from app.services.inspections import (
    apply_checklist,
    build_items,
    load_templates,
    next_inspection_number,
    summarize,
    template_by_slug,
)
from tests.conftest import login

YEAR = date.today().year
SIGNATURE = "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
SMALL_TEMPLATE = {
    "slug": "yard-walk",
    "title": "Yard Walk",
    "kind": "inspection",
    "industry": "general",
    "frequency": "Weekly",
    "intro": "A short walk of the yard.",
    "sections": [
        {"heading": "Yard", "items": [{"text": "Gates close and latch.", "ref": "1910.176"}, {"text": "Salt pile is covered.", "ref": ""}]},
        {"heading": "Fuel island", "items": [{"text": "Spill kit is stocked.", "ref": "1910.120"}]},
    ],
}


def start(client, slug="construction-site-inspection"):
    resp = client.post(f"/inspections/start/{slug}", follow_redirects=True)
    assert resp.status_code == 200
    return resp


def latest_id(app):
    with app.app_context():
        return Inspection.query.order_by(Inspection.id.desc()).first().id


# ---------------------------------------------------------------- templates
def test_templates_load_sorted_and_index_renders_cards(app, client):
    with app.app_context():
        templates = load_templates()
        slugs = [t["slug"] for t in templates]
        assert "construction-site-inspection" in slugs
        assert "general-industry-inspection" in slugs
        assert "annual-safety-program-audit" in slugs
        kinds = [t["kind"] for t in templates]
        assert kinds == sorted(kinds, key=lambda k: 0 if k == "inspection" else 1), "inspections come before audits"
        construction = template_by_slug("construction-site-inspection")
        assert construction["item_count"] == 136 and construction["section_count"] == 15
        assert construction["industry_label"] == "Construction"
        audit = template_by_slug("annual-safety-program-audit")
        assert audit["kind"] == "audit" and audit["item_count"] == 69
        assert template_by_slug("does-not-exist") is None

    login(client, "emp@example.gov")
    resp = client.get("/inspections")
    assert resp.status_code == 200
    assert b"Construction Site Safety Inspection" in resp.data
    assert b"Annual Safety Program Audit" in resp.data
    assert b"Start inspection" in resp.data and b"Start audit" in resp.data
    assert b"136 items in 15 sections" in resp.data
    assert b"/inspections/start/construction-site-inspection" in resp.data
    assert b"No inspections or audits match." in resp.data
    assert b"inspections.css" in resp.data


def test_loader_skips_unreadable_files_and_builds_stable_keys(tmp_path):
    folder = tmp_path / "inspections"
    folder.mkdir()
    (folder / "yard-walk.json").write_text(json.dumps(SMALL_TEMPLATE), encoding="utf-8")
    (folder / "broken.json").write_text('{"slug": "broken", "title": "Half written', encoding="utf-8")
    (folder / "wrong-shape.json").write_text(json.dumps({"slug": "x", "title": "No sections"}), encoding="utf-8")
    templates = load_templates(str(tmp_path))
    assert [t["slug"] for t in templates] == ["yard-walk"]
    items = build_items(templates[0])
    assert [i["key"] for i in items] == ["s0i0", "s0i1", "s1i0"]
    assert items[0] == {"key": "s0i0", "section": "Yard", "text": "Gates close and latch.", "ref": "1910.176", "result": "", "note": ""}
    assert items[2]["section"] == "Fuel island"
    assert load_templates(str(tmp_path / "missing")) == []
    assert not os.path.exists(tmp_path / "missing")


# ---------------------------------------------------------------- starting
def test_start_assigns_numbers_per_kind_and_copies_items(app, client):
    login(client, "emp@example.gov")
    resp = start(client)
    assert f"INS-{YEAR}-001".encode() in resp.data
    resp = start(client)
    assert f"INS-{YEAR}-002".encode() in resp.data
    resp = start(client, "annual-safety-program-audit")
    assert f"AUD-{YEAR}-001".encode() in resp.data
    assert b"Partly in place" in resp.data and b"Missing" in resp.data
    assert b"Department or program audited" in resp.data
    assert client.post("/inspections/start/no-such-template").status_code == 404

    with app.app_context():
        numbers = sorted(i.number for i in Inspection.query.all())
        assert numbers == [f"AUD-{YEAR}-001", f"INS-{YEAR}-001", f"INS-{YEAR}-002"]
        first = Inspection.query.filter_by(number=f"INS-{YEAR}-001").first()
        assert first.kind == "inspection" and first.template_slug == "construction-site-inspection"
        assert first.title == "Construction Site Safety Inspection"
        assert len(first.items) == 136 and first.items[0]["key"] == "s0i0"
        assert all(item["result"] == "" for item in first.items)
        assert first.inspector_name == "James Carter" and first.inspector_id == first.created_by_id
        assert first.division.code == "electric"
        assert first.inspection_date == date.today()
        assert first.completed is False and first.findings == []
        audit = Inspection.query.filter_by(kind="audit").first()
        assert len(audit.items) == 69 and audit.result_labels["partial"] == "Partly in place"
        assert next_inspection_number("inspection") == f"INS-{YEAR}-003"
        assert next_inspection_number("audit") == f"AUD-{YEAR}-002"


# ---------------------------------------------------------------- saving
def test_save_results_notes_findings_signature_and_completion(app, client):
    login(client, "emp@example.gov")
    start(client)
    ins_id = latest_id(app)
    with app.app_context():
        electric_id = db.session.get(Inspection, ins_id).division_id

    header = {
        "site": "Water main replacement, Oak St",
        "inspection_date": "2026-10-20",
        "division_id": str(electric_id),
        "inspector_name": "James Carter",
        "accompanied_by": "Crew leader",
        "weather": "Clear, 60 F",
        "summary": "Housekeeping needs attention near the trench.",
    }
    answers = {"r_s0i0": "ok", "r_s0i1": "action", "n_s0i1": "Form lumber with nails by the spoil pile", "r_s0i2": "na"}
    resp = client.post(f"/inspections/{ins_id}", data={**header, **answers}, follow_redirects=True)
    assert resp.status_code == 200
    assert b"Saved. 3 of 136 items answered." in resp.data
    # The flagged item now appears in the findings block with its inputs.
    assert b'name="fa_s0i1"' in resp.data and b"Form lumber with nails" in resp.data
    assert b"Action needed" in resp.data

    with app.app_context():
        ins = db.session.get(Inspection, ins_id)
        assert ins.site == "Water main replacement, Oak St"
        assert ins.inspection_date == date(2026, 10, 20)
        assert ins.weather == "Clear, 60 F" and ins.accompanied_by == "Crew leader"
        by_key = {i["key"]: i for i in ins.items}
        assert by_key["s0i0"]["result"] == "ok" and by_key["s0i1"]["result"] == "action" and by_key["s0i2"]["result"] == "na"
        assert by_key["s0i1"]["note"] == "Form lumber with nails by the spoil pile"
        assert by_key["s0i3"]["result"] == ""
        assert ins.findings == [{"key": "s0i1", "action": "", "responsible": "", "target_date": "", "status": "open"}]
        assert ins.completed is False and ins.completed_at is None and ins.signature_data is None
        stats = ins.stats
        assert stats["total"] == 136 and stats["answered"] == 3 and stats["ok"] == 1 and stats["action"] == 1 and stats["na"] == 1
        totals = summarize(ins)
        assert totals["applicable"] == 2 and totals["score"] == 50 and totals["findings"] == 1 and totals["open_findings"] == 1

    # Fill in the finding, sign and mark complete.
    finding = {"fa_s0i1": "Remove lumber and bend or pull nails", "fr_s0i1": "Crew leader", "ft_s0i1": "2026-10-22", "fs_s0i1": "open"}
    resp = client.post(
        f"/inspections/{ins_id}",
        data={**header, **answers, **finding, "signature_data": SIGNATURE, "completed": "yes"},
        follow_redirects=True,
    )
    assert resp.status_code == 200
    assert b"Remove lumber and bend or pull nails" in resp.data
    with app.app_context():
        ins = db.session.get(Inspection, ins_id)
        assert ins.findings[0]["action"] == "Remove lumber and bend or pull nails"
        assert ins.findings[0]["responsible"] == "Crew leader"
        assert ins.findings[0]["target_date"] == "2026-10-22"
        assert ins.signature_data == SIGNATURE
        assert ins.completed is True and ins.completed_at is not None
        completed_at = ins.completed_at
    resp = client.get("/inspections")
    assert b"Completed" in resp.data and b"1 open of 1" in resp.data

    # Saving again without a signature field keeps the signature and the completion time.
    resp = client.post(f"/inspections/{ins_id}", data={**header, **answers, **finding, "fs_s0i1": "closed", "completed": "yes"}, follow_redirects=True)
    assert resp.status_code == 200
    with app.app_context():
        ins = db.session.get(Inspection, ins_id)
        assert ins.signature_data == SIGNATURE and ins.completed_at == completed_at
        assert ins.findings[0]["status"] == "closed" and ins.open_findings == []
    resp = client.get("/inspections")
    assert b"1 closed" in resp.data

    # Changing the result drops the finding; leaving the box unchecked reopens the inspection.
    resp = client.post(f"/inspections/{ins_id}", data={**header, **answers, **finding, "r_s0i1": "ok", "signature_data": ""}, follow_redirects=True)
    assert resp.status_code == 200
    assert b"No findings so far." in resp.data
    with app.app_context():
        ins = db.session.get(Inspection, ins_id)
        assert ins.findings == [] and ins.completed is False and ins.completed_at is None
        assert ins.signature_data is None
        assert ins.stats["answered"] == 3


def test_audit_findings_come_from_partial_and_missing(app, client):
    login(client, "super@example.gov")
    start(client, "annual-safety-program-audit")
    ins_id = latest_id(app)
    resp = client.post(
        f"/inspections/{ins_id}",
        data={"site": "Electric Services", "inspection_date": "2026-10-21", "inspector_name": "Dana Super", "weather": "Calendar year 2026",
              "r_s0i0": "ok", "r_s0i1": "partial", "r_s0i2": "missing", "r_s0i3": "action"},
        follow_redirects=True,
    )
    assert resp.status_code == 200
    with app.app_context():
        ins = db.session.get(Inspection, ins_id)
        assert ins.number.startswith("AUD-")
        assert [f["key"] for f in ins.findings] == ["s0i1", "s0i2"]
        by_key = {i["key"]: i for i in ins.items}
        assert by_key["s0i3"]["result"] == "", "an inspection-only result is not accepted on an audit"
        assert summarize(ins)["answered"] == 3
    assert b"Partly in place" in resp.data and b'name="fa_s0i2"' in resp.data


def test_apply_checklist_keeps_only_findings_for_flagged_items(app):
    with app.app_context():
        ins = Inspection(number="INS-2000-001", template_slug="yard-walk", kind="inspection", title="Yard Walk",
                         items_json=json.dumps(build_items(SMALL_TEMPLATE)),
                         findings_json=json.dumps([{"key": "s1i0", "action": "Restock", "responsible": "Ravi", "target_date": "2026-11-01", "status": "open"}]))
        items, findings = apply_checklist(ins, {"r_s0i0": "action", "r_s1i0": "action", "n_s0i0": "  latch broken  ", "ft_s0i0": "not a date", "fs_s0i0": "weird"})
        assert items[0]["note"] == "latch broken"
        assert findings == [
            {"key": "s0i0", "action": "", "responsible": "", "target_date": "", "status": "open"},
            {"key": "s1i0", "action": "Restock", "responsible": "Ravi", "target_date": "2026-11-01", "status": "open"},
        ]
        items, findings = apply_checklist(ins, {"r_s1i0": "ok"})
        assert findings == [] and items[2]["result"] == "ok" and items[0]["result"] == ""


# ---------------------------------------------------------------- printing
def test_print_page_renders_with_empty_and_filled_data(app, client):
    login(client, "emp@example.gov")
    start(client)
    ins_id = latest_id(app)
    resp = client.get(f"/inspections/{ins_id}/print")
    assert resp.status_code == 200
    assert b'class="print-page"' in resp.data
    assert f"INS-{YEAR}-001".encode() in resp.data
    assert b"Checklist" in resp.data and b"No findings." in resp.data and b"In progress" in resp.data
    assert b"General site and housekeeping" in resp.data

    client.post(
        f"/inspections/{ins_id}",
        data={"site": "Lift station 4", "inspection_date": "2026-10-20", "inspector_name": "James Carter", "summary": "Two items to fix.",
              "r_s0i0": "action", "n_s0i0": "Debris on the stairs", "fa_s0i0": "Clear stairs daily", "fr_s0i0": "Crew leader",
              "ft_s0i0": "2026-10-25", "signature_data": SIGNATURE, "completed": "yes"},
        follow_redirects=True,
    )
    resp = client.get(f"/inspections/{ins_id}/print")
    assert resp.status_code == 200
    assert b"Lift station 4" in resp.data and b"Two items to fix." in resp.data
    assert b"Clear stairs daily" in resp.data and b"Debris on the stairs" in resp.data
    assert SIGNATURE.encode() in resp.data
    assert b"Completed" in resp.data
    assert b"1 of 136 items answered" in resp.data


# ---------------------------------------------------------------- access
def test_access_rules(app, client):
    login(client, "emp@example.gov")
    start(client)
    ins_id = latest_id(app)
    number = f"INS-{YEAR}-001".encode()
    assert client.get(f"/inspections/{ins_id}").status_code == 200
    # The creator cannot delete.
    assert client.post(f"/inspections/{ins_id}/delete").status_code == 403
    client.get("/logout")

    # Another employee in a different division sees nothing of it.
    login(client, "other@example.gov")
    assert client.get(f"/inspections/{ins_id}").status_code == 403
    assert client.post(f"/inspections/{ins_id}", data={"site": "Hijack"}).status_code == 403
    assert client.get(f"/inspections/{ins_id}/print").status_code == 403
    assert client.post(f"/inspections/{ins_id}/delete").status_code == 403
    assert number not in client.get("/inspections").data
    export = load_workbook(BytesIO(client.get("/inspections/export.xlsx").data))
    assert export["Items"].max_row == 1
    client.get("/logout")

    # The supervisor of the division can open and edit it but not delete it.
    with app.app_context():
        electric_id = db.session.get(Inspection, ins_id).division_id
    login(client, "super@example.gov")
    assert number in client.get("/inspections").data
    assert client.get(f"/inspections/{ins_id}").status_code == 200
    resp = client.post(
        f"/inspections/{ins_id}",
        data={"site": "Checked by supervisor", "inspection_date": "2026-10-20", "division_id": str(electric_id), "r_s0i0": "ok"},
        follow_redirects=True,
    )
    assert resp.status_code == 200 and b"Checked by supervisor" in resp.data
    assert client.post(f"/inspections/{ins_id}/delete").status_code == 403

    # Moving it out of the division sends the supervisor back to the list and hides it from them.
    resp = client.post(f"/inspections/{ins_id}", data={"site": "Checked by supervisor", "inspection_date": "2026-10-20", "division_id": "", "r_s0i0": "ok"}, follow_redirects=True)
    assert resp.status_code == 200 and b"outside your division" in resp.data
    assert client.get(f"/inspections/{ins_id}").status_code == 403
    assert number not in client.get("/inspections").data
    client.get("/logout")

    # The creator still has it.
    login(client, "emp@example.gov")
    assert client.get(f"/inspections/{ins_id}").status_code == 200
    client.get("/logout")

    # The administrator sees everything and is the only one who can delete.
    login(client, "admin@example.gov")
    assert number in client.get("/inspections").data
    assert client.get(f"/inspections/{ins_id}").status_code == 200
    assert b"Delete" in client.get(f"/inspections/{ins_id}").data
    resp = client.post(f"/inspections/{ins_id}/delete", follow_redirects=True)
    assert resp.status_code == 200 and b"deleted" in resp.data
    with app.app_context():
        assert db.session.get(Inspection, ins_id) is None
    assert client.get(f"/inspections/{ins_id}").status_code == 404


def test_pages_require_login(client):
    assert client.get("/inspections").status_code == 302
    assert client.post("/inspections/start/construction-site-inspection").status_code == 302
    assert client.get("/inspections/export.xlsx").status_code == 302


# ---------------------------------------------------------------- filters and export
def test_index_filters_by_kind_and_status(app, client):
    login(client, "emp@example.gov")
    start(client)
    start(client, "annual-safety-program-audit")
    audit_id = latest_id(app)
    client.post(f"/inspections/{audit_id}", data={"site": "Department", "inspection_date": "2026-10-21", "completed": "yes"})
    ins = f"INS-{YEAR}-001".encode()
    aud = f"AUD-{YEAR}-001".encode()
    data = client.get("/inspections").data
    assert ins in data and aud in data
    data = client.get("/inspections?kind=audit").data
    assert aud in data and ins not in data
    data = client.get("/inspections?kind=inspection").data
    assert ins in data and aud not in data
    data = client.get("/inspections?status=completed").data
    assert aud in data and ins not in data
    data = client.get("/inspections?status=in_progress").data
    assert ins in data and aud not in data


def test_export_xlsx_has_items_and_findings_sheets(app, client):
    login(client, "emp@example.gov")
    start(client)
    ins_id = latest_id(app)
    client.post(
        f"/inspections/{ins_id}",
        data={"site": "Oak St", "inspection_date": "2026-10-20", "inspector_name": "James Carter",
              "r_s0i0": "ok", "r_s0i1": "action", "n_s0i1": "Nails", "fa_s0i1": "Pull nails", "fr_s0i1": "Crew leader", "ft_s0i1": "2026-10-22"},
    )
    resp = client.get("/inspections/export.xlsx")
    assert resp.status_code == 200
    assert resp.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment" in resp.headers["Content-Disposition"] and ".xlsx" in resp.headers["Content-Disposition"]
    wb = load_workbook(BytesIO(resp.data))
    assert wb.sheetnames == ["Items", "Findings"]
    items = wb["Items"]
    assert [c.value for c in items[1]] == ["Number", "Title", "Kind", "Site", "Date", "Inspector", "Section", "Item", "Reference", "Result", "Note"]
    assert items.max_row == 1 + 136
    first = [c.value for c in items[2]]
    assert first[0] == f"INS-{YEAR}-001" and first[2] == "Site inspection" and first[3] == "Oak St"
    assert first[5] == "James Carter" and first[6] == "General site and housekeeping" and first[8] == "1926.25(a)" and first[9] == "Satisfactory"
    second = [c.value for c in items[3]]
    assert second[9] == "Action needed" and second[10] == "Nails"
    findings = wb["Findings"]
    assert findings.max_row == 2
    row = [c.value for c in findings[2]]
    assert row[0] == f"INS-{YEAR}-001" and row[8] == "Action needed" and row[9] == "Pull nails" and row[10] == "Crew leader"
    assert row[11].date() == date(2026, 10, 22) and row[12] == "Open"
