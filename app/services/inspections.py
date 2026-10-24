"""Inspection helpers: checklist templates, numbering, checklist items, access
control and the Excel export for site inspections and annual audits.

Checklist templates are JSON files in content/inspections/*.json:
  {"slug", "title", "kind", "industry", "frequency", "intro",
   "sections": [{"heading", "items": [{"text", "ref"}]}]}
"""
import glob
import json
import logging
import os
from io import BytesIO

from flask import current_app
from sqlalchemy import or_

from ..models import INSPECTION_KINDS, Inspection
from ..timeutil import now
from ..utils import parse_date

log = logging.getLogger(__name__)

KIND_ORDER = {"inspection": 0, "audit": 1}
INDUSTRY_LABELS = {"construction": "Construction", "general": "General industry", "program": "Safety program"}
FINDING_STATUSES = {"open": "Open", "closed": "Closed"}
# Results that turn an item into a finding with a corrective action.
FINDING_RESULTS = {"inspection": ("action",), "audit": ("partial", "missing")}
STATUS_FILTERS = {"in_progress": "In progress", "completed": "Completed"}


# ---------------------------------------------------------------- templates
def template_files(content_dir=None):
    content_dir = content_dir or current_app.config["CONTENT_DIR"]
    return sorted(glob.glob(os.path.join(content_dir, "inspections", "*.json")))


def _valid_template(data):
    if not isinstance(data, dict) or not data.get("slug") or not data.get("title"):
        return False
    sections = data.get("sections")
    if not isinstance(sections, list) or not sections:
        return False
    for section in sections:
        if not isinstance(section, dict) or not isinstance(section.get("items"), list):
            return False
        if any(not isinstance(item, dict) or not item.get("text") for item in section["items"]):
            return False
    return True


def load_templates(content_dir=None):
    """All readable checklist templates, inspections first, then audits, then by title.
    Files that are missing, half written or malformed are skipped."""
    templates = []
    for path in template_files(content_dir):
        try:
            with open(path, encoding="utf-8") as fh:
                data = json.load(fh)
        except (OSError, ValueError) as exc:
            log.warning("Skipping checklist template %s: %s", path, exc)
            continue
        if not _valid_template(data):
            log.warning("Skipping checklist template %s: unexpected structure", path)
            continue
        if data.get("kind") not in INSPECTION_KINDS:
            data["kind"] = "inspection"
        data.setdefault("industry", "general")
        data.setdefault("frequency", "")
        data.setdefault("intro", "")
        data["industry_label"] = INDUSTRY_LABELS.get(data["industry"], str(data["industry"]).title())
        data["kind_label"] = INSPECTION_KINDS[data["kind"]]
        data["section_count"] = len(data["sections"])
        data["item_count"] = sum(len(s["items"]) for s in data["sections"])
        templates.append(data)
    templates.sort(key=lambda t: (KIND_ORDER.get(t["kind"], 9), t["title"].lower()))
    return templates


def template_by_slug(slug, content_dir=None):
    for template in load_templates(content_dir):
        if template["slug"] == slug:
            return template
    return None


def build_items(template):
    """Flatten the template sections into checklist items with stable keys
    ("s<section index>i<item index>", zero based) and empty results."""
    items = []
    for si, section in enumerate(template.get("sections", [])):
        for ii, item in enumerate(section.get("items", [])):
            items.append({
                "key": f"s{si}i{ii}",
                "section": section.get("heading", ""),
                "text": item.get("text", ""),
                "ref": item.get("ref", "") or "",
                "result": "",
                "note": "",
            })
    return items


# ---------------------------------------------------------------- numbering
def next_inspection_number(kind, when=None):
    """INS-YYYY-NNN for inspections and AUD-YYYY-NNN for audits, counting per year."""
    when = when or now()
    prefix = f"{'AUD' if kind == 'audit' else 'INS'}-{when.year}-"
    highest = 0
    for (number,) in Inspection.query.with_entities(Inspection.number).filter(Inspection.number.like(prefix + "%")).all():
        try:
            highest = max(highest, int(number[len(prefix):]))
        except ValueError:
            continue
    return f"{prefix}{highest + 1:03d}"


# ---------------------------------------------------------------- checklist state
def finding_results(kind):
    return FINDING_RESULTS.get(kind, FINDING_RESULTS["inspection"])


def field_labels(kind):
    """Header wording that differs between a site inspection and a program audit."""
    if kind == "audit":
        return {
            "noun": "audit",
            "site": "Department or program audited",
            "inspector": "Auditor",
            "weather": "Period covered",
            "weather_placeholder": "Calendar year",
            "details": "Audit details",
        }
    return {
        "noun": "inspection",
        "site": "Site or facility",
        "inspector": "Inspector",
        "weather": "Weather and conditions",
        "weather_placeholder": "e.g., clear, 70 F, dry",
        "details": "Inspection details",
    }


def _clean(value):
    return value.strip() if isinstance(value, str) else ""


def _iso_date(value):
    parsed = parse_date(_clean(value))
    return parsed.isoformat() if parsed else ""


def apply_checklist(inspection, form):
    """Copy results, notes and findings from a posted form onto the inspection.

    The form carries the whole checklist: r_<key> radios, n_<key> notes and
    fa_/fr_/ft_/fs_<key> finding fields. An item keeps a finding only while its
    result is one that needs action; a finding whose fields were not posted
    (an item flagged in this very save) keeps or starts an empty record."""
    allowed = inspection.result_labels
    flagged = finding_results(inspection.kind)
    previous = {f.get("key"): f for f in inspection.findings if f.get("key")}
    items = inspection.items
    findings = []
    for item in items:
        key = item.get("key")
        result = form.get(f"r_{key}", "")
        item["result"] = result if result in allowed else ""
        item["note"] = _clean(form.get(f"n_{key}", ""))
        if item["result"] not in flagged:
            continue
        old = previous.get(key, {})
        status = _clean(form.get(f"fs_{key}", old.get("status", "open")))
        findings.append({
            "key": key,
            "action": _clean(form.get(f"fa_{key}", old.get("action", ""))),
            "responsible": _clean(form.get(f"fr_{key}", old.get("responsible", ""))),
            "target_date": _iso_date(form.get(f"ft_{key}", old.get("target_date", ""))),
            "status": status if status in FINDING_STATUSES else "open",
        })
    inspection.items_json = json.dumps(items)
    inspection.findings_json = json.dumps(findings)
    return items, findings


def summarize(inspection):
    """Counts for the page header and the list: answered, score, findings."""
    counts = inspection.stats
    total = counts.get("total", 0)
    answered = counts.get("answered", 0)
    applicable = answered - counts.get("na", 0)
    ok = counts.get("ok", 0)
    flagged = finding_results(inspection.kind)
    findings = sum(1 for item in inspection.items if item.get("result") in flagged)
    return {
        "total": total,
        "answered": answered,
        "applicable": applicable,
        "score": round(100 * ok / applicable) if applicable else 0,
        "percent": round(100 * answered / total) if total else 0,
        "findings": findings,
        "open_findings": len(inspection.open_findings),
        "counts": counts,
    }


def group_sections(items):
    """Items grouped back into their sections, in order, with per-section progress."""
    sections = []
    for item in items:
        heading = item.get("section") or ""
        if not sections or sections[-1]["heading"] != heading:
            # "rows" rather than "items" so Jinja does not pick up dict.items
            sections.append({"heading": heading, "rows": [], "answered": 0, "total": 0})
        section = sections[-1]
        section["rows"].append(item)
        section["total"] += 1
        if item.get("result"):
            section["answered"] += 1
    for section in sections:
        section["percent"] = round(100 * section["answered"] / section["total"]) if section["total"] else 0
    return sections


def findings_with_items(inspection):
    """Each saved finding joined with its checklist item for display."""
    items = {item.get("key"): item for item in inspection.items}
    labels = inspection.result_labels
    rows = []
    for finding in inspection.findings:
        item = items.get(finding.get("key"))
        if item is None:
            continue
        row = dict(item)
        row.update({
            "action": finding.get("action", ""),
            "responsible": finding.get("responsible", ""),
            "target_date": finding.get("target_date", ""),
            "target": parse_date(finding.get("target_date", "")),  # date object for display and export
            "status": finding.get("status") or "open",
            "result_label": labels.get(item.get("result"), ""),
        })
        row["status_label"] = FINDING_STATUSES.get(row["status"], row["status"])
        rows.append(row)
    return rows


# ---------------------------------------------------------------- access
def can_view(user, inspection):
    if user.is_admin:
        return True
    if inspection.created_by_id == user.id or inspection.inspector_id == user.id:
        return True
    if user.is_supervisor and inspection.division_id and inspection.division_id == user.division_id:
        return True
    return False


def can_edit(user, inspection):
    if user.is_admin:
        return True
    if inspection.created_by_id == user.id:
        return True
    return bool(user.is_supervisor and inspection.division_id and inspection.division_id == user.division_id)


def can_delete(user, inspection):
    return bool(user.is_admin)


def visible_inspections_query(user):
    query = Inspection.query
    if user.is_admin:
        return query
    conditions = [Inspection.created_by_id == user.id, Inspection.inspector_id == user.id]
    if user.is_supervisor and user.division_id:
        conditions.append(Inspection.division_id == user.division_id)
    return query.filter(or_(*conditions))


# ---------------------------------------------------------------- export
ITEM_COLUMNS = ["Number", "Title", "Kind", "Site", "Date", "Inspector", "Section", "Item", "Reference", "Result", "Note"]
FINDING_COLUMNS = [
    "Number", "Title", "Kind", "Site", "Date", "Section", "Item", "Reference", "Result",
    "Corrective action", "Responsible", "Target date", "Status",
]


def export_workbook(inspections):
    """One row per checklist item on the first sheet and one per finding on the second."""
    from openpyxl import Workbook

    wb = Workbook()
    items_sheet = wb.active
    items_sheet.title = "Items"
    items_sheet.append(ITEM_COLUMNS)
    findings_sheet = wb.create_sheet("Findings")
    findings_sheet.append(FINDING_COLUMNS)
    for inspection in inspections:
        labels = inspection.result_labels
        base = [inspection.number, inspection.title, inspection.kind_label, inspection.site or "", inspection.inspection_date]
        for item in inspection.items:
            items_sheet.append(base + [
                inspection.inspector_name or "",
                item.get("section", ""),
                item.get("text", ""),
                item.get("ref", ""),
                labels.get(item.get("result"), ""),
                item.get("note", ""),
            ])
        for row in findings_with_items(inspection):
            findings_sheet.append(base + [
                row.get("section", ""),
                row.get("text", ""),
                row.get("ref", ""),
                row.get("result_label", ""),
                row.get("action", ""),
                row.get("responsible", ""),
                row.get("target") or "",
                row.get("status_label", ""),
            ])
    for sheet in (items_sheet, findings_sheet):
        sheet.freeze_panes = "A2"
        for column in sheet.iter_cols(min_row=1, max_row=1):
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = 42 if column[0].value in ("Item", "Corrective action") else 16
    return wb


def export_bytes(inspections):
    buffer = BytesIO()
    export_workbook(inspections).save(buffer)
    buffer.seek(0)
    return buffer
