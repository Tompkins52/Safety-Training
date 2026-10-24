"""Excel exports of training records, built with openpyxl.

One workbook layout serves three audiences: the administrator (every active
employee, optionally one division), a supervisor (direct reports) and an
employee (their own record). The caller decides which users go in; nothing in
here widens that list.

Sheets:
  Training record   one row per assignment
  Quiz attempts     one row per quiz attempt
  Compliance        one row per employee, computed from the assignment rows
  Courses           the course catalog
"""
import io
from datetime import date, datetime

from flask import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..models import Assignment, Course, QuizAttempt

XLSX_MIMETYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DATE_FORMAT = "yyyy-mm-dd"
DATETIME_FORMAT = "yyyy-mm-dd hh:mm"

RECORD_HEADERS = [
    "Employee", "Employee number", "Division", "Job title", "Supervisor", "Course", "Standard", "Category",
    "Reason", "Assigned on", "Due date", "Status", "Completed on", "Best score", "Attempts", "Certificate number",
]
ATTEMPT_HEADERS = ["Employee", "Course", "Taken at", "Score", "Total", "Percent", "Passed"]
COMPLIANCE_HEADERS = [
    "Employee", "Employee number", "Division", "Supervisor", "Assigned", "Completed", "Overdue",
    "Due within 30 days", "Percent complete",
]
COURSE_HEADERS = ["Slug", "Title", "Standard", "Category", "New hire", "Divisions", "Renewal months", "Active"]

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="173A5E")
MIN_WIDTH, MAX_WIDTH = 10, 48


def _sort_key(user):
    return ((user.last_name or "").lower(), (user.first_name or "").lower(), user.id)


def _write_sheet(ws, headers, rows):
    """Bold header row, frozen top row, real date cells, widths from the content."""
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    widths = [len(h) for h in headers]
    for row in rows:
        ws.append(row)
        for index, value in enumerate(row):
            if value is None:
                continue
            if isinstance(value, datetime):
                ws.cell(row=ws.max_row, column=index + 1).number_format = DATETIME_FORMAT
                length = 16
            elif isinstance(value, date):
                ws.cell(row=ws.max_row, column=index + 1).number_format = DATE_FORMAT
                length = 10
            else:
                length = len(str(value))
            if length > widths[index]:
                widths[index] = length
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = max(MIN_WIDTH, min(MAX_WIDTH, width + 2))
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions


def record_rows(users, assignments_by_user):
    rows = []
    for user in users:
        for a in assignments_by_user.get(user.id, []):
            course = a.course
            rows.append([
                user.full_name,
                user.employee_number or "",
                user.division.name if user.division else "",
                user.job_title or "",
                user.supervisor.full_name if user.supervisor else "",
                course.title,
                course.standard or "",
                course.category_label,
                a.reason_label,
                a.assigned_on,
                a.due_date,
                a.display_status,
                a.completed_at if a.is_completed else None,
                a.best_score if a.best_score is not None else None,
                a.attempts or 0,
                a.certificate_number if a.is_completed else "",
            ])
    return rows


def attempt_rows(users, attempts_by_user):
    rows = []
    for user in users:
        for attempt in attempts_by_user.get(user.id, []):
            rows.append([
                user.full_name,
                attempt.course.title,
                attempt.taken_at,
                attempt.score,
                attempt.total,
                attempt.percent,
                "Yes" if attempt.passed else "No",
            ])
    return rows


def compliance_rows(users, assignments_by_user, today=None):
    """Per employee: assigned, completed, overdue, due within 30 days, percent complete.

    Counts are over the assignment rows in the Training record sheet. Each
    course counts once, by its latest assignment: completed when that
    assignment is completed, or when it is the automatic renewal of an
    earlier completion and is not yet overdue (the employee is current on
    the course). Overdue and due within 30 days count open assignments.
    """
    today = today or date.today()
    rows = []
    for user in users:
        latest, ever_completed = {}, set()
        for a in assignments_by_user.get(user.id, []):
            current = latest.get(a.course_id)
            if current is None or (a.cycle, a.id) > (current.cycle, current.id):
                latest[a.course_id] = a
            if a.is_completed:
                ever_completed.add(a.course_id)
        assigned = len(latest)
        completed = overdue = due_soon = 0
        for course_id, a in latest.items():
            if a.is_completed:
                completed += 1
                continue
            days_left = (a.due_date - today).days
            if days_left < 0:
                overdue += 1
            else:
                if days_left <= 30:
                    due_soon += 1
                if course_id in ever_completed:
                    completed += 1
        rows.append([
            user.full_name,
            user.employee_number or "",
            user.division.name if user.division else "",
            user.supervisor.full_name if user.supervisor else "",
            assigned,
            completed,
            overdue,
            due_soon,
            round(100 * completed / assigned) if assigned else None,
        ])
    return rows


def course_rows(courses):
    return [
        [
            c.slug,
            c.title,
            c.standard or "",
            c.category_label,
            "Yes" if c.new_hire else "No",
            ", ".join(d.name for d in c.divisions),
            c.renewal_months,
            "Yes" if c.is_active else "No",
        ]
        for c in courses
    ]


def training_workbook(users, today=None):
    """Build the training record workbook for exactly these users."""
    today = today or date.today()
    users = sorted(users, key=_sort_key)
    user_ids = [u.id for u in users]
    assignments_by_user, attempts_by_user = {}, {}
    if user_ids:
        assignments = (
            Assignment.query.filter(Assignment.user_id.in_(user_ids))
            .order_by(Assignment.due_date, Assignment.cycle, Assignment.id)
            .all()
        )
        for a in assignments:
            assignments_by_user.setdefault(a.user_id, []).append(a)
        attempts = QuizAttempt.query.filter(QuizAttempt.user_id.in_(user_ids)).order_by(QuizAttempt.taken_at, QuizAttempt.id).all()
        for attempt in attempts:
            attempts_by_user.setdefault(attempt.user_id, []).append(attempt)
    courses = Course.query.order_by(Course.category, Course.top10_rank, Course.title).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Training record"
    _write_sheet(ws, RECORD_HEADERS, record_rows(users, assignments_by_user))
    _write_sheet(wb.create_sheet("Quiz attempts"), ATTEMPT_HEADERS, attempt_rows(users, attempts_by_user))
    _write_sheet(wb.create_sheet("Compliance"), COMPLIANCE_HEADERS, compliance_rows(users, assignments_by_user, today))
    _write_sheet(wb.create_sheet("Courses"), COURSE_HEADERS, course_rows(courses))
    return wb


def workbook_bytes(wb):
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def workbook_response(wb, stem="training-record", today=None):
    """Flask response that downloads the workbook as <stem>-YYYY-MM-DD.xlsx."""
    today = today or date.today()
    filename = f"{stem}-{today:%Y-%m-%d}.xlsx"
    return Response(
        workbook_bytes(wb),
        mimetype=XLSX_MIMETYPE,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def training_export_response(users, stem="training-record"):
    """Shortcut used by the three export routes."""
    return workbook_response(training_workbook(users), stem=stem)
