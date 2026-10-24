"""Supervisor reminder copies, the new hire track, individual assignments and Excel exports."""
import io
from datetime import date, timedelta

from openpyxl import load_workbook

from app.extensions import db
from app.models import Assignment, Course, Division, NotificationLog, User
from app.services.assignments import (
    assign_individual,
    assign_new_hire_training,
    ensure_assignments,
    new_hire_courses,
    new_hire_progress,
    run_annual_assignment,
)
from app.services.exports import ATTEMPT_HEADERS, COMPLIANCE_HEADERS, COURSE_HEADERS, RECORD_HEADERS, XLSX_MIMETYPE
from app.services.notifications import run_daily_notifications
from tests.conftest import login

SUPERVISOR_KINDS = ("training_due_30_supervisor", "training_due_7_supervisor", "training_overdue_supervisor")


def _user(email):
    return User.query.filter_by(email=email).first()


def _course(slug):
    return Course.query.filter_by(slug=slug).first()


def _log(kind):
    return NotificationLog.query.filter_by(kind=kind).order_by(NotificationLog.id).all()


# --------------------------------------------------------------------------- A. supervisor copies


def test_supervisor_gets_copies_of_reminders_without_duplicates(app):
    with app.app_context():
        emp = _user("emp@example.gov")
        today = date.today()
        a30 = ensure_assignments(emp, due_date=today + timedelta(days=30), courses=[_course("ladders")])[0]
        a7 = ensure_assignments(emp, due_date=today + timedelta(days=5), courses=[_course("scaffolding")])[0]
        aover = ensure_assignments(emp, due_date=today - timedelta(days=3), courses=[_course("hazard-communication")])[0]
        db.session.commit()

        counts = run_daily_notifications()
        assert counts["training_due_30"] == 1 and counts["training_due_7"] == 1 and counts["training_overdue"] == 1
        assert counts["training_due_30_supervisor"] == 1
        assert counts["training_due_7_supervisor"] == 1
        assert counts["training_overdue_supervisor"] == 1

        for kind, assignment in ((SUPERVISOR_KINDS[0], a30), (SUPERVISOR_KINDS[1], a7), (SUPERVISOR_KINDS[2], aover)):
            entries = _log(kind)
            assert len(entries) == 1, kind
            entry = entries[0]
            assert entry.recipient_email == "super@example.gov"
            assert entry.recipient_user_id == _user("super@example.gov").id
            assert entry.assignment_id == assignment.id
            assert "James Carter" in entry.body
            assert assignment.course.title in entry.body
            assert assignment.due_date.strftime("%B %d, %Y") in entry.body
            assert "/team" in entry.body
        assert "OVERDUE" in _log(SUPERVISOR_KINDS[2])[0].subject
        assert "3 days overdue" in _log(SUPERVISOR_KINDS[2])[0].body
        # Employee reminders still go to the employee, not the supervisor.
        assert all(e.recipient_email == "emp@example.gov" for e in _log("training_due_30") + _log("training_due_7") + _log("training_overdue"))

        # A second run the same day sends nothing new to anyone.
        counts = run_daily_notifications()
        assert counts["training_due_30"] == 0 and counts["training_due_7"] == 0 and counts["training_overdue"] == 0
        assert all(counts[k] == 0 for k in SUPERVISOR_KINDS)
        assert counts["skipped"] == 3
        assert sum(len(_log(k)) for k in SUPERVISOR_KINDS) == 3
        assert NotificationLog.query.count() == 6

        # The admin notification log page lists the new kinds with their labels.
    client = app.test_client()
    login(client, "admin@example.gov")
    resp = client.get("/admin/notifications")
    assert resp.status_code == 200
    assert b"30-day reminder (supervisor copy)" in resp.data
    assert b"7-day reminder (supervisor copy)" in resp.data
    assert b"Overdue reminder (supervisor copy)" in resp.data
    resp = client.get("/admin/notifications?kind=training_overdue_supervisor")
    assert resp.status_code == 200
    assert b"super@example.gov" in resp.data


def test_supervisor_copy_is_one_digest_per_employee_and_kind(app):
    with app.app_context():
        emp = _user("emp@example.gov")
        today = date.today()
        # Three courses coming due within the 30-day window on the same run: one digest listing all three.
        courses = [_course("ladders"), _course("scaffolding"), _course("machine-guarding")]
        for i, course in enumerate(courses):
            ensure_assignments(emp, due_date=today + timedelta(days=20 + i), courses=[course])
        # Two overdue courses: one overdue digest.
        for course in (_course("hazard-communication"), _course("lockout-tagout")):
            ensure_assignments(emp, due_date=today - timedelta(days=10), courses=[course])
        db.session.commit()

        counts = run_daily_notifications()
        assert counts["training_due_30"] == 3 and counts["training_overdue"] == 2
        assert counts["training_due_30_supervisor"] == 1 and counts["training_overdue_supervisor"] == 1
        due_digest = _log("training_due_30_supervisor")
        assert len(due_digest) == 1
        for course in courses:
            assert course.title in due_digest[0].body
        assert "3 training modules coming due" in due_digest[0].subject
        over_digest = _log("training_overdue_supervisor")
        assert len(over_digest) == 1
        assert "Hazard Communication" in over_digest[0].body and "Lockout/Tagout" in over_digest[0].body
        assert "2 training modules past due" in over_digest[0].subject

        # Still nothing new on a second run.
        counts = run_daily_notifications()
        assert counts["training_due_30_supervisor"] == 0 and counts["training_overdue_supervisor"] == 0


def test_no_supervisor_means_no_copy_and_no_error(app):
    with app.app_context():
        other = _user("other@example.gov")  # no supervisor on file
        assert other.supervisor is None
        ensure_assignments(other, due_date=date.today() - timedelta(days=1), courses=[_course("ladders")])
        db.session.commit()
        counts = run_daily_notifications()
        assert counts["training_overdue"] == 1
        assert counts["training_overdue_supervisor"] == 0
        assert _log("training_overdue_supervisor") == []
        assert len(_log("training_overdue")) == 1


# --------------------------------------------------------------------------- B. new hire track


def test_assign_new_hire_training_is_limited_to_new_hire_courses_and_idempotent(app):
    with app.app_context():
        emp = _user("emp@example.gov")
        other = _user("other@example.gov")
        flagged = {c.slug for c in Course.query.filter_by(new_hire=True).all()}
        assert len(flagged) == 10
        assert "lockout-tagout" in flagged and "pole-top-rescue" not in flagged and "incident-reporting-investigation" not in flagged
        assert {c.slug for c in new_hire_courses(emp)} == flagged

        # No hire date: due DEFAULT_DUE_DAYS from today.
        created = assign_new_hire_training(emp)
        db.session.commit()
        assert len(created) == 10
        assert all(a.reason == "new_hire" for a in created)
        assert {a.course.slug for a in created} == flagged
        assert all(a.due_date == date.today() + timedelta(days=app.config["DEFAULT_DUE_DAYS"]) for a in created)
        assert all(a.course.new_hire for a in emp.assignments)

        # Idempotent.
        assert assign_new_hire_training(emp) == []
        assert len(emp.assignments) == 10

        # Hire date: due hire date + 30 days; assigned_by recorded.
        other.hire_date = date.today() - timedelta(days=10)
        admin = _user("admin@example.gov")
        created = assign_new_hire_training(other, assigned_by=admin, send_notice=True)
        db.session.commit()
        assert len(created) == 10
        assert all(a.due_date == other.hire_date + timedelta(days=30) for a in created)
        assert all(a.assigned_by_id == admin.id for a in created)
        notices = _log("training_assigned")
        assert len(notices) == 10 and all(n.recipient_email == "other@example.gov" for n in notices)
        assert "new hire training track" in notices[0].body

        progress = new_hire_progress(other)
        assert progress["total"] == 10 and progress["completed"] == 0 and progress["percent"] == 0
        assert progress["next_due"] == other.hire_date + timedelta(days=30)

        # The annual pass adds only the remaining courses, with reason annual, and does not duplicate.
        created = run_annual_assignment(due_date=date.today() + timedelta(days=60))
        by_user = {}
        for a in created:
            by_user.setdefault(a.user.email, []).append(a)
        assert len(by_user["emp@example.gov"]) == 18 - 10
        assert len(by_user["other@example.gov"]) == 16 - 10
        assert all(a.reason == "annual" for a in created)
        assert len(emp.open_assignments()) == 18
        assert len({a.course_id for a in emp.assignments}) == 18

        # A user with no division gets nothing.
        assert assign_new_hire_training(admin) == []


def test_new_hire_renews_as_annual_after_completion(app):
    with app.app_context():
        emp = _user("emp@example.gov")
        assignment = assign_new_hire_training(emp)[0]
        db.session.commit()
        from app.services.assignments import grade_quiz

        right = {q.id: q.answer_index for q in assignment.course.questions}
        grade_quiz(assignment, right)
        assert assignment.is_completed
        renewal = Assignment.query.filter_by(user_id=emp.id, course_id=assignment.course_id, cycle=2).first()
        assert renewal is not None and renewal.reason == "annual"
        progress = new_hire_progress(emp)
        assert progress["completed"] == 1 and progress["total"] == 10 and progress["percent"] == 10


def test_admin_new_hires_page_and_assign_button(app, client):
    with app.app_context():
        emp = _user("emp@example.gov")
        emp.hire_date = date.today() - timedelta(days=20)
        other = _user("other@example.gov")
        other.hire_date = date.today() - timedelta(days=200)  # not a recent hire, no new hire assignments
        db.session.commit()
        emp_id = emp.id

    login(client, "admin@example.gov")
    resp = client.get("/admin/new-hires")
    assert resp.status_code == 200
    assert b"New hire training" in resp.data
    assert b"James Carter" in resp.data
    assert b"Ravi Patel" not in resp.data
    assert b"Not assigned" in resp.data
    assert b"Required new hire courses (10)" in resp.data

    resp = client.post(f"/admin/employees/{emp_id}/new-hire", data={"next": "/admin/new-hires"}, follow_redirects=True)
    assert resp.status_code == 200
    assert b"New hire training assigned to James Carter: 10 modules" in resp.data
    assert b"0 of 10 complete" in resp.data
    with app.app_context():
        emp = db.session.get(User, emp_id)
        items = [a for a in emp.assignments if a.reason == "new_hire"]
        assert len(items) == 10
        assert all(a.due_date == emp.hire_date + timedelta(days=30) for a in items)
        assert all(a.assigned_by.email == "admin@example.gov" for a in items)

    # Second click: nothing duplicated.
    resp = client.post(f"/admin/employees/{emp_id}/new-hire", data={}, follow_redirects=True)
    assert b"already has an open assignment for every new hire course" in resp.data
    with app.app_context():
        assert Assignment.query.filter_by(user_id=emp_id).count() == 10

    # The employee edit page groups the record by reason.
    resp = client.get(f"/admin/employees/{emp_id}/edit")
    assert resp.status_code == 200
    assert b"Assign new hire training" in resp.data
    assert b"0 of 10 complete" in resp.data
    assert resp.data.count(b'<h3 class="section-label">New hire') == 1

    # An employee without a division cannot get the track.
    with app.app_context():
        admin_id = _user("admin@example.gov").id
    resp = client.post(f"/admin/employees/{admin_id}/new-hire", data={}, follow_redirects=True)
    assert b"has no division" in resp.data


def test_employee_create_form_assigns_new_hire_track(app, client):
    login(client, "admin@example.gov")
    with app.app_context():
        public_id = Division.query.filter_by(code="public").first().id
    hire_date = date.today() - timedelta(days=3)
    resp = client.post(
        "/admin/employees/new",
        data={
            "first_name": "Pat",
            "last_name": "Newcomer",
            "email": "pat.newcomer@example.gov",
            "role": "employee",
            "division_id": str(public_id),
            "hire_date": hire_date.strftime("%Y-%m-%d"),
            "is_active": "on",
            "new_hire_now": "on",
            "assign_now": "on",
            "due_date": (date.today() + timedelta(days=90)).strftime("%Y-%m-%d"),
        },
        follow_redirects=True,
    )
    assert resp.status_code == 200
    assert b"New hire training assigned to Pat Newcomer: 10 modules" in resp.data
    with app.app_context():
        user = _user("pat.newcomer@example.gov")
        new_hire = [a for a in user.assignments if a.reason == "new_hire"]
        annual = [a for a in user.assignments if a.reason == "annual"]
        assert len(new_hire) == 10
        assert all(a.due_date == hire_date + timedelta(days=30) for a in new_hire)
        assert len(annual) == 18 - 10  # public services: rescue courses apply, supervisor course does not
        assert all(a.due_date == date.today() + timedelta(days=90) for a in annual)
        user_id = user.id

    resp = client.get("/admin/employees")
    assert resp.status_code == 200
    assert b"Pat Newcomer" in resp.data
    resp = client.get(f"/admin/employees/{user_id}/edit")
    assert b"Assigned individually" not in resp.data and b"Annual" in resp.data


# --------------------------------------------------------------------------- C. individual assignments


def test_individual_assignment_bypasses_division_rule_and_never_duplicates(app, client):
    with app.app_context():
        other = _user("other@example.gov")  # Environmental: rescue and supervisor courses do not apply
        rescue = _course("pole-top-rescue")
        supervisor_course = _course("incident-reporting-investigation")
        assert not rescue.applies_to(other) and not supervisor_course.applies_to(other)
        other_id, rescue_id, sup_id = other.id, rescue.id, supervisor_course.id
        ladders_id = _course("ladders").id
        ensure_assignments(other, due_date=date.today() + timedelta(days=40), courses=[_course("ladders")])
        db.session.commit()

    login(client, "admin@example.gov")
    resp = client.get(f"/admin/assign/individual?user={other_id}")
    assert resp.status_code == 200
    assert b"Assign individual training" in resp.data
    assert b"Supervisors and managers" in resp.data
    due = date.today() + timedelta(days=21)
    resp = client.post(
        "/admin/assign/individual",
        data={
            "user_id": str(other_id),
            "courses": [str(rescue_id), str(sup_id), str(ladders_id)],
            "due_date": due.strftime("%Y-%m-%d"),
            "note": "Acting supervisor during the November rotation.",
            "notify": "on",
        },
        follow_redirects=True,
    )
    assert resp.status_code == 200
    assert b"2 courses assigned to Ravi Patel" in resp.data
    assert b"not duplicated: Ladders" in resp.data
    with app.app_context():
        individual = Assignment.query.filter_by(user_id=other_id, reason="individual").all()
        assert {a.course_id for a in individual} == {rescue_id, sup_id}
        assert all(a.due_date == due and a.assigned_by.email == "admin@example.gov" for a in individual)
        assert Assignment.query.filter_by(user_id=other_id, course_id=ladders_id).count() == 1
        notices = _log("training_assigned")
        assert len(notices) == 2
        assert all("Acting supervisor during the November rotation." in n.body for n in notices)
        assert all("assigned to you individually" in n.body for n in notices)

    # Posting the same courses again creates nothing.
    resp = client.post(
        "/admin/assign/individual",
        data={"user_id": str(other_id), "courses": [str(rescue_id), str(sup_id)], "due_date": due.strftime("%Y-%m-%d")},
        follow_redirects=True,
    )
    assert b"not duplicated" in resp.data
    with app.app_context():
        assert Assignment.query.filter_by(user_id=other_id).count() == 3
        # The service reports the skip too.
        other = db.session.get(User, other_id)
        created, skipped = assign_individual(other, [db.session.get(Course, rescue_id)], due)
        assert created == [] and [c.id for c in skipped] == [rescue_id]

    # Validation.
    resp = client.post("/admin/assign/individual", data={"user_id": "", "due_date": ""}, follow_redirects=True)
    assert b"Choose an employee." in resp.data and b"Tick at least one course." in resp.data and b"Enter a valid due date." in resp.data

    # The reason shows as a tag on My Training.
    client.get("/logout")
    login(client, "other@example.gov")
    resp = client.get("/training")
    assert resp.status_code == 200
    assert resp.data.count(b"Assigned individually") == 2
    assert b"Pole Top Rescue" in resp.data


def test_individual_course_outside_division_is_not_renewed(app):
    with app.app_context():
        other = _user("other@example.gov")
        rescue = _course("pole-top-rescue")
        created, _skipped = assign_individual(other, [rescue], date.today() + timedelta(days=30), send_notice=False)
        db.session.commit()
        from app.services.assignments import grade_quiz

        assignment = created[0]
        grade_quiz(assignment, {q.id: q.answer_index for q in rescue.questions})
        assert assignment.is_completed
        assert Assignment.query.filter_by(user_id=other.id, course_id=rescue.id).count() == 1


def test_team_page_shows_reason_tags(app, client):
    with app.app_context():
        emp = _user("emp@example.gov")
        assign_new_hire_training(emp)
        assign_individual(emp, [_course("incident-reporting-investigation")], date.today() + timedelta(days=15), send_notice=False)
        db.session.commit()
    login(client, "super@example.gov")
    resp = client.get("/team")
    assert resp.status_code == 200
    assert b"0 of 10" in resp.data
    assert b"Assigned individually" in resp.data
    assert resp.data.count(b'>New hire</span>') == 10
    assert b"Export to Excel" in resp.data


# --------------------------------------------------------------------------- D. Excel exports


def _open(resp):
    assert resp.status_code == 200
    assert resp.mimetype == XLSX_MIMETYPE
    assert resp.headers["Content-Disposition"].endswith(f'{date.today():%Y-%m-%d}.xlsx"')
    return load_workbook(io.BytesIO(resp.data))


def _rows(ws):
    return [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]


def _headers(ws):
    return [c.value for c in ws[1]]


def _seed_training(app):
    with app.app_context():
        run_annual_assignment(due_date=date.today() + timedelta(days=20))
        emp = _user("emp@example.gov")
        from app.services.assignments import grade_quiz

        assignment = next(a for a in emp.assignments if a.course.slug == "ladders")
        questions = list(assignment.course.questions)
        grade_quiz(assignment, {q.id: (q.answer_index + 1) % 4 for q in questions})  # fail
        grade_quiz(assignment, {q.id: q.answer_index for q in questions})  # pass
        other = _user("other@example.gov")
        overdue = next(a for a in other.assignments if a.course.slug == "ladders")
        overdue.due_date = date.today() - timedelta(days=2)
        db.session.commit()


def test_admin_excel_export(app, client):
    _seed_training(app)
    login(client, "admin@example.gov")
    resp = client.get("/admin/reports")
    assert b"Export to Excel" in resp.data
    wb = _open(client.get("/admin/reports/export.xlsx"))
    assert wb.sheetnames == ["Training record", "Quiz attempts", "Compliance", "Courses"]

    record = wb["Training record"]
    assert _headers(record) == RECORD_HEADERS
    assert record.freeze_panes == "A2"
    assert all(c.font.bold for c in record[1])
    rows = _rows(record)
    employees = {r[0] for r in rows}
    assert employees == {"James Carter", "Dana Super", "Ravi Patel"}
    carter = [r for r in rows if r[0] == "James Carter"]
    assert len(carter) == 19  # 18 annual plus the renewal created when Ladders was passed
    completed = next(r for r in carter if r[5] == "Ladders" and r[11] == "Completed")
    assert completed[2] == "Electric Services" and completed[4] == "Dana Super"
    assert completed[7] == "OSHA Top 10" and completed[8] == "Annual"
    assert isinstance(completed[9], date) and isinstance(completed[10], date)
    assert completed[12] is not None and completed[13] == 10 and completed[14] == 2
    assert completed[15].startswith("CERT-")
    assert record.cell(row=2, column=10).number_format == "yyyy-mm-dd"
    assert any(r[11] == "Overdue" and r[0] == "Ravi Patel" for r in rows)
    for index in range(1, len(RECORD_HEADERS) + 1):
        assert record.column_dimensions[record.cell(row=1, column=index).column_letter].width >= 10

    attempts = wb["Quiz attempts"]
    assert _headers(attempts) == ATTEMPT_HEADERS
    arows = _rows(attempts)
    assert len(arows) == 2
    assert arows[0][:2] == ["James Carter", "Ladders"] and arows[0][6] == "No" and arows[1][6] == "Yes"
    assert arows[1][3] == 10 and arows[1][4] == 10 and arows[1][5] == 100

    compliance = wb["Compliance"]
    assert _headers(compliance) == COMPLIANCE_HEADERS
    crows = {r[0]: r for r in _rows(compliance)}
    # Every active employee in scope appears, including one with nothing assigned.
    assert set(crows) == employees | {"Sam Admin"}
    assert crows["Sam Admin"][4] == 0 and crows["Sam Admin"][8] is None
    carter_row = crows["James Carter"]
    assert carter_row[4] == 18 and carter_row[5] == 1 and carter_row[6] == 0 and carter_row[7] == 17 and carter_row[8] == 6
    patel_row = crows["Ravi Patel"]
    assert patel_row[4] == 16 and patel_row[6] == 1 and patel_row[7] == 15 and patel_row[8] == 0

    courses = wb["Courses"]
    assert _headers(courses)[:7] == COURSE_HEADERS[:7]
    course_rows = _rows(courses)
    assert len(course_rows) == 19
    ladders = next(r for r in course_rows if r[0] == "ladders")
    assert ladders[1] == "Ladders" and ladders[4] == "Yes" and ladders[6] == 12
    assert "Electric Services" in ladders[5]

    # Division filter narrows the workbook.
    with app.app_context():
        electric_id = Division.query.filter_by(code="electric").first().id
    wb = _open(client.get(f"/admin/reports/export.xlsx?division={electric_id}"))
    assert {r[0] for r in _rows(wb["Training record"])} == {"James Carter", "Dana Super"}


def test_employee_excel_export_has_only_their_rows(app, client):
    _seed_training(app)
    login(client, "emp@example.gov")
    resp = client.get("/training")
    assert b"Export to Excel" in resp.data
    wb = _open(client.get("/training/export.xlsx"))
    assert wb.sheetnames == ["Training record", "Quiz attempts", "Compliance", "Courses"]
    for sheet in ("Training record", "Quiz attempts", "Compliance"):
        rows = _rows(wb[sheet])
        assert rows, sheet
        assert {r[0] for r in rows} == {"James Carter"}, sheet
    assert len(_rows(wb["Training record"])) == 19
    assert len(_rows(wb["Quiz attempts"])) == 2
    # Admin routes stay closed to employees.
    assert client.get("/admin/reports/export.xlsx").status_code == 403
    assert client.get("/team/export.xlsx").status_code == 403


def test_supervisor_team_export_is_limited_to_direct_reports(app, client):
    _seed_training(app)
    login(client, "super@example.gov")
    wb = _open(client.get("/team/export.xlsx"))
    rows = _rows(wb["Training record"])
    assert {r[0] for r in rows} == {"James Carter"}
    assert "Ravi Patel" not in {r[0] for r in _rows(wb["Compliance"])}
    assert len(_rows(wb["Courses"])) == 19


# --------------------------------------------------------------------------- E. dashboard, catalog, about


def test_dashboard_shows_new_hire_block(app, client):
    with app.app_context():
        emp = _user("emp@example.gov")
        emp.hire_date = date.today() - timedelta(days=5)
        assign_new_hire_training(emp)
        # The other employee gets ordinary annual assignments only.
        ensure_assignments(_user("other@example.gov"), due_date=date.today() + timedelta(days=40), courses=[_course("ladders")])
        db.session.commit()

    login(client, "other@example.gov")
    resp = client.get("/")
    assert resp.status_code == 200
    assert b"New hire training" not in resp.data
    client.get("/logout")

    login(client, "emp@example.gov")
    resp = client.get("/")
    assert resp.status_code == 200
    assert b"New hire training" in resp.data
    assert b"0 of 10 complete" in resp.data
    assert b"and 4 more" in resp.data


def test_catalog_and_about_mention_new_features(client):
    login(client, "emp@example.gov")
    resp = client.get("/catalog")
    assert resp.status_code == 200
    assert resp.data.count(b'<span class="tag">New hire</span>') == 10
    assert b"Supervisors and managers" in resp.data
    resp = client.get("/about")
    assert resp.status_code == 200
    for phrase in (b"new hire training track", b"Individual assignments", b"Excel export", b"direct supervisor", b"Eight-section incident reports", b"Site inspections and annual program audits"):
        assert phrase in resp.data, phrase


# --------------------------------------------------------------------------- F. CLI


def test_cli_assign_new_hire_and_seed_demo(app):
    runner = app.test_cli_runner()
    result = runner.invoke(args=["assign-new-hire", "emp@example.gov"])
    assert result.exit_code == 0, result.output
    assert "10 new hire assignments created for James Carter" in result.output
    result = runner.invoke(args=["assign-new-hire", "emp@example.gov"])
    assert "already has an open assignment" in result.output
    result = runner.invoke(args=["assign-new-hire", "nobody@example.gov"])
    assert result.exit_code != 0
    result = runner.invoke(args=["assign-new-hire", "admin@example.gov"])
    assert result.exit_code != 0 and "no division" in result.output

    result = runner.invoke(args=["seed-demo"])
    assert result.exit_code == 0, result.output
    assert "New hire:" in result.output
    with app.app_context():
        mai = _user("m.nguyen@example.gov")
        assert mai.hire_date is not None and 0 <= (date.today() - mai.hire_date).days <= 30
        new_hire = [a for a in mai.assignments if a.reason == "new_hire"]
        assert len(new_hire) == 10
        assert all(a.due_date == mai.hire_date + timedelta(days=30) for a in new_hire)
        assert len([a for a in mai.assignments if a.reason == "annual"]) == 8
    result = runner.invoke(args=["assign-annual"])
    assert result.exit_code == 0, result.output
